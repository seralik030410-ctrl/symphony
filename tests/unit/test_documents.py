import copy
import json
from pathlib import Path

import pytest
from pydantic import ValidationError

from backend.artifacts.schemas import EXAMPLES, parse_spec
from backend.artifacts.formulas import calculate
from backend.artifacts.renderers import render_job, render_docx, render_slides


@pytest.mark.parametrize("format", ["pdf", "xlsx", "docx", "pptx"])
def test_examples_validate_and_reject_arbitrary_renderer_code(format):
    assert parse_spec(format, EXAMPLES[format]).title
    with pytest.raises(ValidationError): parse_spec(format, {**EXAMPLES[format], "python": "import os"})


def test_formulas_calculated_and_cross_sheet_references():
    value = copy.deepcopy(EXAMPLES["xlsx"])
    value["sheets"][0]["formulas"]["B3"] = "=ROUND(SUM(B2:B2)/3,2)"
    value["sheets"].append({"name": "Other", "columns": [{"name": "Value", "type": "number"}], "rows": [[None]], "formulas": {"A2": "='Бюджет'!B3*2"}})
    result = calculate(parse_spec("xlsx", value))
    assert result["values"]["Бюджет!B3"] == 33.33
    assert result["values"]["Other!A2"] == 66.66


@pytest.mark.parametrize("formula", ["=HYPERLINK(\"https://evil\",\"x\")", "=WEBSERVICE(\"https://evil\")", "='[secret.xlsx]Sheet1'!A1", "=B3", "=B9000", "=SUM(A1:A1)", "=1/0", "=__import__('os')", "=CMD|' /C calc'!A0"])
def test_unsafe_circular_and_invalid_formulas_are_rejected(formula):
    value = copy.deepcopy(EXAMPLES["xlsx"]); value["sheets"][0]["formulas"]["B3"] = formula
    with pytest.raises((ValueError, ZeroDivisionError, SyntaxError)): calculate(parse_spec("xlsx", value))


@pytest.mark.parametrize("change", ["wrongtype", "ragged", "duplicate", "nan"])
def test_workbook_types_and_shape(change):
    value = copy.deepcopy(EXAMPLES["xlsx"])
    if change == "wrongtype": value["sheets"][0]["rows"][0][1] = "100"
    elif change == "ragged": value["sheets"][0]["rows"][0] = []
    elif change == "duplicate": value["sheets"].append(copy.deepcopy(value["sheets"][0]))
    else: value["sheets"][0]["rows"][0][1] = float("nan")
    with pytest.raises(ValidationError): parse_spec("xlsx", value)


@pytest.mark.parametrize("format", ["pdf", "xlsx"])
def test_actual_renderer_output_and_preview(tmp_path, format):
    value = copy.deepcopy(EXAMPLES[format])
    if format == "pdf":
        value["sections"][0]["paragraphs"] = ["Без исполнения <script>alert(1)</script>. Кириллица и проверенные данные."]
        value["sections"][0]["table"] = {"columns": ["Статья", "Сумма"], "rows": [["Факт", 120], ["План", 100]]}
        value["sections"][0]["chart"] = {"title": "План и факт", "labels": ["План", "Факт"], "values": [100, 120]}
    else: value["sheets"][0]["rows"][0][0] = "=HYPERLINK(\"https://evil\")"
    (tmp_path / "source.json").write_text(json.dumps(value), encoding="utf-8")
    (tmp_path / "recipe.json").write_text(json.dumps({"format": format}), encoding="utf-8")
    result = render_job(tmp_path)
    assert result["valid"] and (tmp_path / f"document.{format}").stat().st_size > 1000
    if format == "pdf":
        import pymupdf
        with pymupdf.open(tmp_path / "document.pdf") as document:
            text = "".join(page.get_text() for page in document)
        assert "Кириллица" in text and "<script>" in text
        assert (tmp_path / "page-001.png").stat().st_size > 1000
    else:
        from openpyxl import load_workbook
        book = load_workbook(tmp_path / "document.xlsx", data_only=False)
        assert book.active["A2"].data_type == "s"
        assert book.active["B3"].data_type == "f"
        assert result["tables"][0]["rows"][1][1] == 100


def test_office_renderers_create_reopenable_native_files(tmp_path):
    from docx import Document
    from pptx import Presentation
    render_docx(parse_spec("docx", EXAMPLES["docx"]), tmp_path / "document.docx")
    render_slides(parse_spec("pptx", EXAMPLES["pptx"]), tmp_path / "document.pptx")
    assert Document(tmp_path / "document.docx").paragraphs[0].text == "Документ"
    assert len(Presentation(tmp_path / "document.pptx").slides) == 1
