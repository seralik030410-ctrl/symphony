"""Advanced spreadsheet analysis and formula audit engine for FinControl IDE."""
from __future__ import annotations

from pathlib import Path
from typing import Any
import re
import math

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import pandas as pd
except ImportError:
    pd = None


FORMULA_ERROR_PATTERNS = [
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
]


def analyze_spreadsheet(
    file_path: Path | str,
    sheet_name: str | None = None,
    deep: bool = True,
) -> dict[str, Any]:
    """Analyzes an Excel spreadsheet (.xlsx/.xlsm/.csv), computing metrics, data profiles, and formula audit."""
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        return _analyze_csv(path)
    elif suffix in (".xlsx", ".xlsm", ".xltx"):
        return _analyze_xlsx(path, sheet_name=sheet_name, deep=deep)
    else:
        raise ValueError(f"Unsupported spreadsheet format for analysis: {suffix}")


def _analyze_xlsx(path: Path, sheet_name: str | None = None, deep: bool = True) -> dict[str, Any]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed")

    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheet_names = wb.sheetnames
    target_sheet_name = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]
    sheet = wb[target_sheet_name]

    max_row = sheet.max_row or 0
    max_col = sheet.max_column or 0

    # Load evaluated data version to compare formulas and calculated values
    try:
        data_wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        data_sheet = data_wb[target_sheet_name] if target_sheet_name in data_wb.sheetnames else None
    except Exception:
        data_sheet = None

    rows_formula: list[list[Any]] = []
    rows_values: list[list[Any]] = []
    formulas_found: list[dict[str, Any]] = []
    formula_errors: list[dict[str, Any]] = []

    for r in range(1, max_row + 1):
        r_formulas: list[Any] = []
        r_vals: list[Any] = []
        for c in range(1, max_col + 1):
            cell_formula = sheet.cell(row=r, column=c)
            f_val = cell_formula.value
            r_formulas.append(f_val)

            val = f_val
            if data_sheet is not None:
                try:
                    val = data_sheet.cell(row=r, column=c).value
                except Exception:
                    pass
            r_vals.append(val)

            coord = cell_formula.coordinate
            if isinstance(f_val, str) and f_val.startswith("="):
                formula_record = {
                    "cell": coord,
                    "row": r,
                    "col": c,
                    "formula": f_val,
                    "evaluated_value": str(val) if val is not None else "",
                }
                formulas_found.append(formula_record)

                val_str = str(val) if val is not None else ""
                for err in FORMULA_ERROR_PATTERNS:
                    if err in val_str or err in f_val:
                        formula_errors.append({
                            "cell": coord,
                            "error": err,
                            "formula": f_val,
                        })
                        break
            elif isinstance(val, str):
                for err in FORMULA_ERROR_PATTERNS:
                    if err in val:
                        formula_errors.append({
                            "cell": coord,
                            "error": err,
                            "formula": "",
                        })
                        break

        rows_formula.append(r_formulas)
        rows_values.append(r_vals)

    # Headers & Column profiling
    headers: list[str] = []
    if rows_values:
        first_row = rows_values[0]
        for idx, h in enumerate(first_row):
            h_str = str(h).strip() if h is not None else ""
            headers.append(h_str if h_str else f"Col_{openpyxl.utils.get_column_letter(idx + 1)}")
    else:
        headers = [f"Col_{openpyxl.utils.get_column_letter(i + 1)}" for i in range(max_col)]

    data_rows = rows_values[1:] if len(rows_values) > 1 else rows_values

    # Pandas deep dataframe analytics if available
    column_profiles: list[dict[str, Any]] = []
    correlations: dict[str, dict[str, float]] = {}
    summary_stats: dict[str, Any] = {}

    if pd is not None and data_rows:
        try:
            df = pd.DataFrame(data_rows, columns=headers[:len(data_rows[0])] if data_rows else None)
            # Try converting numeric columns
            for col_name in df.columns:
                series = df[col_name]
                numeric_series = pd.to_numeric(series, errors="coerce")
                non_null_numeric = numeric_series.dropna()

                total_count = len(series)
                null_count = int(series.isna().sum()) + int((series == "").sum())
                unique_vals = int(series.nunique())

                # If majority numeric, profile as numeric
                if len(non_null_numeric) > 0 and len(non_null_numeric) >= (total_count - null_count) * 0.7:
                    col_type = "numeric"
                    stats = {
                        "min": float(non_null_numeric.min()),
                        "max": float(non_null_numeric.max()),
                        "mean": round(float(non_null_numeric.mean()), 4),
                        "median": round(float(non_null_numeric.median()), 4),
                        "sum": round(float(non_null_numeric.sum()), 4),
                        "std": round(float(non_null_numeric.std()), 4) if len(non_null_numeric) > 1 else 0.0,
                    }
                    df[col_name] = numeric_series
                else:
                    col_type = "text"
                    stats = None

                column_profiles.append({
                    "name": str(col_name),
                    "type": col_type,
                    "total_count": total_count,
                    "null_count": null_count,
                    "unique_count": unique_vals,
                    "stats": stats,
                    "sample_values": [str(x) for x in series.dropna().unique()[:5] if str(x).strip()],
                })

            # Numeric correlations if >= 2 numeric columns
            numeric_cols = df.select_dtypes(include=["number"]).columns.tolist()
            if len(numeric_cols) >= 2 and deep:
                corr_matrix = df[numeric_cols].corr()
                for c1 in numeric_cols:
                    correlations[c1] = {}
                    for c2 in numeric_cols:
                        val = corr_matrix.loc[c1, c2]
                        if not math.isnan(val):
                            correlations[c1][c2] = round(float(val), 3)

            summary_stats = {
                "total_rows": len(df),
                "total_columns": len(df.columns),
                "numeric_columns_count": len(numeric_cols),
                "text_columns_count": len(df.columns) - len(numeric_cols),
            }
        except Exception as exc:
            summary_stats = {"error": f"Pandas analysis fallback: {exc}"}
    else:
        # Pure Python fallback
        for c_idx in range(max_col):
            col_vals = [r[c_idx] for r in data_rows if c_idx < len(r) and r[c_idx] not in (None, "")]
            nums: list[float] = []
            for v in col_vals:
                try:
                    nums.append(float(v))
                except (ValueError, TypeError):
                    pass

            is_numeric = len(nums) > 0 and len(nums) >= len(col_vals) * 0.7
            col_name = headers[c_idx] if c_idx < len(headers) else f"Col_{c_idx + 1}"
            stats = None
            if is_numeric and nums:
                stats = {
                    "min": min(nums),
                    "max": max(nums),
                    "mean": round(sum(nums) / len(nums), 4),
                    "sum": round(sum(nums), 4),
                }

            column_profiles.append({
                "name": col_name,
                "type": "numeric" if is_numeric else "text",
                "total_count": len(data_rows),
                "null_count": len(data_rows) - len(col_vals),
                "unique_count": len(set(col_vals)),
                "stats": stats,
                "sample_values": [str(x) for x in list(set(col_vals))[:5]],
            })

        summary_stats = {
            "total_rows": len(data_rows),
            "total_columns": max_col,
        }

    return {
        "file_name": path.name,
        "sheet_name": target_sheet_name,
        "available_sheets": sheet_names,
        "dimensions": {
            "max_row": max_row,
            "max_column": max_col,
            "data_row_count": len(data_rows),
        },
        "summary": summary_stats,
        "headers": headers,
        "columns": column_profiles,
        "correlations": correlations,
        "formula_audit": {
            "total_formulas": len(formulas_found),
            "error_count": len(formula_errors),
            "errors": formula_errors[:50],
            "formulas_sample": formulas_found[:25],
        },
    }


def _analyze_csv(path: Path) -> dict[str, Any]:
    """Analyzes a CSV file using pandas or pure python."""
    if pd is not None:
        try:
            df = pd.read_csv(path)
            column_profiles: list[dict[str, Any]] = []
            for col in df.columns:
                series = df[col]
                num_series = pd.to_numeric(series, errors="coerce")
                valid_num = num_series.dropna()
                is_num = len(valid_num) > 0 and len(valid_num) >= len(series.dropna()) * 0.7
                stats = None
                if is_num:
                    stats = {
                        "min": float(valid_num.min()),
                        "max": float(valid_num.max()),
                        "mean": round(float(valid_num.mean()), 4),
                        "sum": round(float(valid_num.sum()), 4),
                    }
                column_profiles.append({
                    "name": str(col),
                    "type": "numeric" if is_num else "text",
                    "total_count": len(series),
                    "null_count": int(series.isna().sum()),
                    "unique_count": int(series.nunique()),
                    "stats": stats,
                    "sample_values": [str(x) for x in series.dropna().unique()[:5]],
                })

            return {
                "file_name": path.name,
                "sheet_name": "CSV",
                "available_sheets": ["CSV"],
                "dimensions": {"max_row": len(df), "max_column": len(df.columns)},
                "summary": {"total_rows": len(df), "total_columns": len(df.columns)},
                "headers": list(df.columns),
                "columns": column_profiles,
                "formula_audit": {"total_formulas": 0, "error_count": 0, "errors": []},
            }
        except Exception:
            pass

    import csv
    with path.open(encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        rows = list(reader)

    headers = rows[0] if rows else []
    return {
        "file_name": path.name,
        "sheet_name": "CSV",
        "available_sheets": ["CSV"],
        "dimensions": {"max_row": len(rows), "max_column": len(headers)},
        "summary": {"total_rows": len(rows), "total_columns": len(headers)},
        "headers": headers,
        "columns": [{"name": h, "type": "text"} for h in headers],
        "formula_audit": {"total_formulas": 0, "error_count": 0, "errors": []},
    }
